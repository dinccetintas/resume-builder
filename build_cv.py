#!/usr/bin/env python3
"""Render Dinc Cetintas' confirmed AI Engineer CV to ATS-safe DOCX + PDF (+ Markdown).

Single source of truth: the CV data below. Both renderers consume the same
structure so the DOCX and PDF stay identical. No tables/text-boxes/images in the
DOCX so applicant tracking systems parse it as clean linear text.
"""
from __future__ import annotations

import os

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor
from fpdf import FPDF

OUT_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), "out")
BASENAME = "Dinc_Cetintas_AI_Engineer"

# --- Confirmed CV content ---------------------------------------------------

NAME = "DINC CETINTAS"
ROLE = "AI Engineer"
CONTACT = (
    "EU Citizen — no visa sponsorship required  ·  Cyprus  ·  "
    "dinccetintas24@gmail.com  ·  +974 51444964  ·  "
    "linkedin.com/in/dinc-cetintas-79a04b205  ·  GitHub"
)

PROFILE = (
    "AI Engineer specializing in Large Language Models, Generative AI, and production ML "
    "systems. I design, fine-tune, and deploy transformer-based models (GPT, T5, BERT) and "
    "RAG pipelines for business-critical use cases — document intelligence, conversational "
    "assistants, and workflow automation. Proven track record of shipping AI to production "
    "with MLOps (Docker, Kubernetes, MLflow) and optimizing models for real-time, low-latency "
    "inference. Strong bias for action and a record of AI solutions that drive measurable "
    "business impact."
)

EXPERIENCE = [
    {
        "title": "AI Engineer / Data Scientist — Commercial Bank of Qatar, Doha",
        "dates": "Jan 2024 – Present",
        "bullets": [
            "Shipped production LLM systems (GPT, T5, BERT) for personalization, document "
            "summarization, and workflow automation, cutting manual processing across multiple "
            "business units by ~60%, by fine-tuning transformer models with Hugging Face "
            "Transformers and orchestrating them through LangChain.",
            "Delivered real-time answers from 900+ unstructured compliance documents at 87% "
            "source-linked retrieval accuracy by building a RAG assistant with an embedding + "
            "vector-retrieval pipeline (semantic chunking, re-ranking, citation grounding).",
            "Reduced manual workload 70% and earned the COO Problem Solver Award by building a "
            "generative-AI email assistant that auto-classifies and responds to 500+ customer "
            "queries weekly (intent classification + LLM-generated drafts).",
            "Extracted structured data from complex Arabic/English ID and financial documents "
            "— cutting document-handling time ~65% — by building 7 OCR + translation "
            "pipelines on Azure OCR and transformer models.",
            "Flagged anomalous transactions and reduced fraud exposure by developing a "
            "supervised fraud-detection model (XGBoost, Logistic Regression) with advanced "
            "feature engineering, validated via A/B testing and ROC-AUC 0.93, handling severe "
            "class imbalance with resampling.",
            "Eliminated third-party translation spend (~$120K/yr) by launching an on-prem "
            "Arabic↔English document translator that preserves original layout, deployed "
            "fully in-house for data privacy.",
            "Cut production inference latency ~40% and enabled reproducible CI/CD model releases "
            "by containerizing and automating LLM deployment with Docker, MLflow, and Kubernetes.",
        ],
    },
    {
        "title": "Software Engineer — MeeApps, İstanbul",
        "dates": "Feb 2023 – Jan 2024",
        "bullets": [
            "Delivered microsecond-level trade execution for 23 brokerage firms managing $10M+ "
            "portfolios by building a high-frequency trading platform and optimizing ITCH, OUCH, "
            "and TCP protocols for ultra-low latency.",
            "Built the full platform — frontend UI, backend services, and live market-data "
            "feeds from co-located exchange servers — using React, C#, and C++.",
            "Owned end-to-end system reliability and latency across the full stack, sustaining "
            "99.9% uptime under live trading load.",
        ],
    },
    {
        "title": "Data Scientist — Schneider Electric, İstanbul",
        "dates": "Jul 2022 – Dec 2022",
        "bullets": [
            "Reduced inventory gaps 12–15% during COVID-related supply shocks by developing "
            "and validating demand-forecasting and stock-out-risk models (Random Forest, XGBoost).",
            "Improved forecast accuracy by engineering features from sales, supply-chain, and "
            "distributor data and applying time-series analysis and model tuning.",
            "Enabled scalable, reproducible deployment by designing evaluation frameworks "
            "(cross-validation, holdout testing) and building model pipelines.",
            "Aligned model outputs with logistics and fulfillment planning by partnering with "
            "stakeholders on success metrics; delivered analysis via SQL, Python (pandas, "
            "scikit-learn), and Power BI.",
        ],
    },
]

EDUCATION = [
    ("BASc, Electronics Engineering — Sabancı University, İstanbul", "2018 – 2023"),
    ("Exchange Semester, Electrical & Electronics Engineering — Universitat "
     "Politècnica de València (UPV), Spain", "2022"),
]

SKILLS = [
    ("Generative AI / LLMs", "RAG, fine-tuning, prompt engineering, transformers (GPT, BERT, T5), "
     "Hugging Face, LangChain, OpenAI APIs, embeddings & vector retrieval"),
    ("ML / Data Science", "Supervised & unsupervised learning, XGBoost, Random Forest, feature "
     "engineering, time-series forecasting, statistical analysis & hypothesis testing, A/B "
     "testing, model evaluation"),
    ("MLOps & Deployment", "Docker, Kubernetes, MLflow, CI/CD, real-time inference optimization, "
     "model serving"),
    ("Cloud", "Azure (OCR, Translation, Custom Models), AWS"),
    ("Programming", "Python (pandas, scikit-learn), SQL; Full-stack (React, C#, C++)"),
]

COURSES = (
    "Machine Learning Specialization (Stanford) · Visa Data Science Bootcamp · "
    "TensorFlow in AI/ML/DL (DeepLearning.AI) · Operationalizing LLMs on Azure (Duke) · "
    "Generative AI Applications with RAG & LangChain (IBM) · AI Agents (IBM) · "
    "Generative AI Techniques & Applications (Duke) · Deep Learning for Sequences & Time "
    "Series (DeepLearning.AI) · Supervised/Unsupervised Learning with scikit-learn (DataCamp)"
)

ADDITIONAL = "Languages: English  ·  Links: LinkedIn, GitHub"


# --- DOCX renderer ----------------------------------------------------------

def build_docx(path: str) -> None:
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10.5)

    for section in doc.sections:
        section.top_margin = section.bottom_margin = Pt(40)
        section.left_margin = section.right_margin = Pt(50)

    def heading(text: str) -> None:
        p = doc.add_paragraph()
        p.space_before = Pt(8)
        run = p.add_run(text.upper())
        run.bold = True
        run.font.size = Pt(11)
        run.font.color.rgb = RGBColor(0x1F, 0x3A, 0x5F)
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(2)

    def bullet(text: str) -> None:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(3)
        p.add_run(text)

    # Header
    name_p = doc.add_paragraph()
    name_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = name_p.add_run(NAME)
    r.bold = True
    r.font.size = Pt(20)
    name_p.paragraph_format.space_after = Pt(0)

    role_p = doc.add_paragraph()
    role_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rr = role_p.add_run(ROLE)
    rr.font.size = Pt(12)
    rr.font.color.rgb = RGBColor(0x1F, 0x3A, 0x5F)
    role_p.paragraph_format.space_after = Pt(2)

    contact_p = doc.add_paragraph()
    contact_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cr = contact_p.add_run(CONTACT)
    cr.font.size = Pt(8.5)
    contact_p.paragraph_format.space_after = Pt(4)

    # Profile
    heading("Profile")
    doc.add_paragraph(PROFILE)

    # Experience
    heading("Professional Experience")
    for job in EXPERIENCE:
        jp = doc.add_paragraph()
        jp.paragraph_format.space_before = Pt(6)
        jp.paragraph_format.space_after = Pt(1)
        tr = jp.add_run(job["title"])
        tr.bold = True
        tr.font.size = Pt(10.5)
        jp.add_run("\t")
        dr = jp.add_run(job["dates"])
        dr.italic = True
        dr.font.size = Pt(9.5)
        for b in job["bullets"]:
            bullet(b)

    # Education
    heading("Education")
    for school, dates in EDUCATION:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        sr = p.add_run(school)
        sr.bold = True
        p.add_run("  —  " + dates).italic = True

    # Skills
    heading("Technical Skills")
    for cat, items in SKILLS:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        p.add_run(cat + ": ").bold = True
        p.add_run(items)

    # Courses
    heading("Certifications & Courses")
    doc.add_paragraph(COURSES)

    # Additional
    heading("Additional")
    doc.add_paragraph(ADDITIONAL)

    doc.save(path)


# --- PDF renderer -----------------------------------------------------------

NAVY = (31, 58, 95)


class CVPdf(FPDF):
    pass


def build_pdf(path: str) -> None:
    pdf = CVPdf(format="A4", unit="mm")
    pdf.set_auto_page_break(auto=True, margin=14)
    pdf.add_page()
    pdf.set_margins(left=16, top=14, right=16)
    width = pdf.w - 32

    def name_block() -> None:
        pdf.set_font("Helvetica", "B", 22)
        pdf.cell(0, 9, NAME, align="C", new_x="LMARGIN", new_y="NEXT")
        pdf.set_text_color(*NAVY)
        pdf.set_font("Helvetica", "", 12)
        pdf.cell(0, 6, ROLE, align="C", new_x="LMARGIN", new_y="NEXT")
        pdf.set_text_color(0, 0, 0)
        pdf.set_font("Helvetica", "", 8)
        pdf.multi_cell(0, 4, latin(CONTACT), align="C")
        pdf.ln(1.5)

    def heading(text: str) -> None:
        pdf.ln(1.5)
        pdf.set_text_color(*NAVY)
        pdf.set_font("Helvetica", "B", 11)
        pdf.cell(0, 5.5, text.upper(), new_x="LMARGIN", new_y="NEXT")
        y = pdf.get_y()
        pdf.set_draw_color(*NAVY)
        pdf.line(16, y, 16 + width, y)
        pdf.ln(1)
        pdf.set_text_color(0, 0, 0)

    def body(text: str) -> None:
        pdf.set_x(pdf.l_margin)
        pdf.set_font("Helvetica", "", 10)
        pdf.multi_cell(width, 4.6, latin(text))

    def bullet(text: str) -> None:
        pdf.set_x(pdf.l_margin)
        pdf.set_font("Helvetica", "", 10)
        pdf.cell(4, 4.6, latin("-"))
        pdf.multi_cell(width - 4, 4.6, latin(text))

    name_block()

    heading("Profile")
    body(PROFILE)

    heading("Professional Experience")
    for job in EXPERIENCE:
        pdf.ln(1.5)
        # Title (bold) and dates (italic, right-aligned) on one row.
        pdf.set_x(pdf.l_margin)
        pdf.set_font("Helvetica", "I", 9)
        date_w = pdf.get_string_width(latin(job["dates"])) + 2
        pdf.set_font("Helvetica", "B", 10.5)
        pdf.cell(width - date_w, 5, latin(job["title"]))
        pdf.set_font("Helvetica", "I", 9)
        pdf.cell(date_w, 5, latin(job["dates"]), align="R", new_x="LMARGIN", new_y="NEXT")
        pdf.ln(0.5)
        for b in job["bullets"]:
            bullet(b)

    heading("Education")
    for school, dates in EDUCATION:
        pdf.set_x(pdf.l_margin)
        pdf.set_font("Helvetica", "B", 10)
        body_str = f"{school}  -  {dates}"
        pdf.multi_cell(width, 4.6, latin(body_str))

    heading("Technical Skills")
    for cat, items in SKILLS:
        pdf.set_font("Helvetica", "B", 10)
        pdf.write(4.8, latin(cat + ": "))
        pdf.set_font("Helvetica", "", 10)
        pdf.write(4.8, latin(items))
        pdf.ln(5.2)

    heading("Certifications & Courses")
    body(COURSES)

    heading("Additional")
    body(ADDITIONAL)

    pdf.output(path)


def latin(s: str) -> str:
    """fpdf2 core fonts are latin-1; map the few unicode glyphs we use."""
    repl = {
        "—": "-", "–": "-", "•": "-", "·": "|",
        "↔": "<->", "İ": "I", "ı": "i", "ç": "c", "Ç": "C",
        "ş": "s", "Ş": "S", "ğ": "g", "Ğ": "G",
        "è": "e", "í": "i", "ü": "u", "ö": "o",
        "‘": "'", "’": "'", "“": '"', "”": '"',
    }
    for k, v in repl.items():
        s = s.replace(k, v)
    return s.encode("latin-1", "replace").decode("latin-1")


# --- Markdown source --------------------------------------------------------

def build_md(path: str) -> None:
    lines = [f"# {NAME}", f"**{ROLE}**", "", CONTACT, "", "## PROFILE", PROFILE, "",
             "## PROFESSIONAL EXPERIENCE"]
    for job in EXPERIENCE:
        lines += ["", f"**{job['title']}** · {job['dates']}"]
        lines += [f"- {b}" for b in job["bullets"]]
    lines += ["", "## EDUCATION"]
    lines += [f"**{s}** · {d}  " for s, d in EDUCATION]
    lines += ["", "## TECHNICAL SKILLS"]
    lines += [f"- **{c}:** {i}" for c, i in SKILLS]
    lines += ["", "## CERTIFICATIONS & COURSES", COURSES, "", "## ADDITIONAL", ADDITIONAL, ""]
    with open(path, "w", encoding="utf-8") as f:
        f.write("\n".join(lines))


def main() -> None:
    os.makedirs(OUT_DIR, exist_ok=True)
    docx_path = os.path.join(OUT_DIR, BASENAME + ".docx")
    pdf_path = os.path.join(OUT_DIR, BASENAME + ".pdf")
    md_path = os.path.join(os.path.dirname(OUT_DIR), BASENAME + ".md")
    build_docx(docx_path)
    build_pdf(pdf_path)
    build_md(md_path)
    for p in (docx_path, pdf_path, md_path):
        print(f"wrote {p}  ({os.path.getsize(p)} bytes)")


if __name__ == "__main__":
    main()
