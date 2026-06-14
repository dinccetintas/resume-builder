"""ATS-safe DOCX / PDF / Markdown renderers, parameterized by ``CVData``.

This is the proven rendering from the original ``build_cv.py``, refactored to
take any ``CVData`` instance so tailored (per-job) CVs render identically to the
default one. No tables / text-boxes / images in the DOCX so ATS parses it as
clean linear text.
"""
from __future__ import annotations

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor
from fpdf import FPDF

from .data import CVData, DEFAULT_CV

NAVY_HEX = "1F3A5F"
INK_HEX = "1A1A1A"
MUTED_HEX = "565656"
NAVY = (31, 58, 95)


# --- DOCX helpers -----------------------------------------------------------

def _set_char_spacing(run, pts: float) -> None:
    rpr = run._element.get_or_add_rPr()
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:val"), str(int(pts * 20)))
    rpr.append(spacing)


def _add_bottom_border(paragraph, color: str = NAVY_HEX, sz: int = 6, space: int = 3) -> None:
    ppr = paragraph._p.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(sz))
    bottom.set(qn("w:space"), str(space))
    bottom.set(qn("w:color"), color)
    pbdr.append(bottom)
    ppr.append(pbdr)


def _add_hyperlink(paragraph, text: str, url: str, *, size: float = 8.5,
                   color: str = NAVY_HEX) -> None:
    r_id = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    col = OxmlElement("w:color")
    col.set(qn("w:val"), color)
    rpr.append(col)
    szel = OxmlElement("w:sz")
    szel.set(qn("w:val"), str(int(size * 2)))
    rpr.append(szel)
    run.append(rpr)
    t = OxmlElement("w:t")
    t.set(qn("xml:space"), "preserve")
    t.text = text
    run.append(t)
    link.append(run)
    paragraph._p.append(link)


# --- DOCX renderer ----------------------------------------------------------

def render_docx(cv: CVData, path: str) -> None:
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10.5)

    sec = doc.sections[0]
    sec.top_margin = sec.bottom_margin = Pt(38)
    sec.left_margin = sec.right_margin = Pt(54)
    right_pos = sec.page_width - sec.left_margin - sec.right_margin

    def heading(text: str) -> None:
        p = doc.add_paragraph()
        run = p.add_run(text.upper())
        run.bold = True
        run.font.size = Pt(10.5)
        run.font.name = "Calibri"
        run.font.color.rgb = RGBColor(0x1F, 0x3A, 0x5F)
        _set_char_spacing(run, 1.2)
        p.paragraph_format.space_before = Pt(11)
        p.paragraph_format.space_after = Pt(4)
        _add_bottom_border(p)

    def bullet(text: str) -> None:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(3)
        p.add_run(text)

    name_p = doc.add_paragraph()
    name_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = name_p.add_run(cv.name)
    r.bold = True
    r.font.size = Pt(24)
    r.font.name = "Georgia"
    r.font.color.rgb = RGBColor(0x1A, 0x1A, 0x1A)
    _set_char_spacing(r, 1.0)
    name_p.paragraph_format.space_after = Pt(1)

    role_p = doc.add_paragraph()
    role_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rr = role_p.add_run(cv.role.upper())
    rr.font.size = Pt(11)
    rr.font.color.rgb = RGBColor(0x1F, 0x3A, 0x5F)
    _set_char_spacing(rr, 2.4)
    role_p.paragraph_format.space_after = Pt(3)

    contact_p = doc.add_paragraph()
    contact_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sep = "   |   "
    for i, (text, url) in enumerate(cv.contact_segments):
        if url:
            _add_hyperlink(contact_p, text, url, size=8.5)
        else:
            run = contact_p.add_run(text)
            run.font.size = Pt(8.5)
            run.font.color.rgb = RGBColor(0x56, 0x56, 0x56)
        if i < len(cv.contact_segments) - 1:
            s = contact_p.add_run(sep)
            s.font.size = Pt(8.5)
            s.font.color.rgb = RGBColor(0xAA, 0xAA, 0xAA)
    contact_p.paragraph_format.space_after = Pt(9)
    _add_bottom_border(contact_p, sz=4, space=4)

    heading("Profile")
    doc.add_paragraph(cv.profile)

    heading("Professional Experience")
    for job in cv.experience:
        jp = doc.add_paragraph()
        jp.paragraph_format.space_before = Pt(7)
        jp.paragraph_format.space_after = Pt(1)
        jp.paragraph_format.tab_stops.add_tab_stop(right_pos, WD_TAB_ALIGNMENT.RIGHT)
        tr = jp.add_run(job.title)
        tr.bold = True
        tr.font.size = Pt(10.5)
        jp.add_run("\t")
        dr = jp.add_run(job.dates)
        dr.italic = True
        dr.font.size = Pt(9.5)
        dr.font.color.rgb = RGBColor(0x56, 0x56, 0x56)
        for b in job.bullets:
            bullet(b)

    heading("Education")
    for school, dates in cv.education:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        sr = p.add_run(school)
        sr.bold = True
        p.add_run("  —  " + dates).italic = True

    heading("Technical Skills")
    for cat, items in cv.skills:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        p.add_run(cat + ": ").bold = True
        p.add_run(items)

    heading("Selected AI Systems")
    for name, desc in cv.selected_systems:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(2)
        p.add_run(name + " — ").bold = True
        p.add_run(desc)

    heading("Certifications & Courses")
    doc.add_paragraph(cv.courses)

    heading("Additional")
    doc.add_paragraph(cv.additional)

    doc.save(path)


# --- PDF renderer -----------------------------------------------------------

def _latin(s: str) -> str:
    """fpdf2 core fonts are latin-1; map the few unicode glyphs we use."""
    repl = {
        "—": "-", "–": "-", "•": "-", "·": "|",
        "↔": "<->", "İ": "I", "ı": "i", "ç": "c", "Ç": "C",
        "ş": "s", "Ş": "S", "ğ": "g", "Ğ": "G",
        "è": "e", "í": "i", "ü": "u", "ö": "o",
        "‘": "'", "’": "'", "“": '"', "”": '"',
    }
    for k, v in repl.items():
        s = s.replace(k, v)
    return s.encode("latin-1", "replace").decode("latin-1")


def render_pdf(cv: CVData, path: str) -> None:
    pdf = FPDF(format="A4", unit="mm")
    pdf.set_auto_page_break(auto=True, margin=14)
    pdf.add_page()
    pdf.set_margins(left=16, top=14, right=16)
    width = pdf.w - 32

    def render_contact_line(size: float = 8.5) -> None:
        pdf.set_font("Helvetica", "", size)
        sep = "   |   "
        sep_w = pdf.get_string_width(sep)
        widths = [pdf.get_string_width(_latin(t)) for t, _ in cv.contact_segments]
        total = sum(widths) + sep_w * (len(cv.contact_segments) - 1)
        pdf.set_x((pdf.w - total) / 2)
        for i, (text, url) in enumerate(cv.contact_segments):
            if url:
                pdf.set_text_color(*NAVY)
                pdf.cell(widths[i], 5, _latin(text), link=url)
            else:
                pdf.set_text_color(0x56, 0x56, 0x56)
                pdf.cell(widths[i], 5, _latin(text))
            if i < len(cv.contact_segments) - 1:
                pdf.set_text_color(0xAA, 0xAA, 0xAA)
                pdf.cell(sep_w, 5, sep)
        pdf.set_text_color(0, 0, 0)
        pdf.ln(6)

    def name_block() -> None:
        pdf.set_font("Times", "B", 24)
        pdf.set_text_color(0x1A, 0x1A, 0x1A)
        pdf.cell(0, 10, _latin(cv.name), align="C", new_x="LMARGIN", new_y="NEXT")
        pdf.set_text_color(*NAVY)
        pdf.set_font("Helvetica", "", 10.5)
        pdf.cell(0, 6, _latin(cv.role.upper()), align="C", new_x="LMARGIN", new_y="NEXT")
        pdf.set_text_color(0, 0, 0)
        render_contact_line()
        y = pdf.get_y() + 0.5
        pdf.set_draw_color(*NAVY)
        pdf.set_line_width(0.4)
        pdf.line(16, y, 16 + width, y)
        pdf.ln(3)

    def heading(text: str) -> None:
        pdf.ln(1.5)
        pdf.set_text_color(*NAVY)
        pdf.set_font("Helvetica", "B", 11)
        pdf.cell(0, 5.5, text.upper(), new_x="LMARGIN", new_y="NEXT")
        y = pdf.get_y()
        pdf.set_draw_color(*NAVY)
        pdf.line(16, y, 16 + width, y)
        pdf.ln(1)
        pdf.set_text_color(0, 0, 0)

    def body(text: str) -> None:
        pdf.set_x(pdf.l_margin)
        pdf.set_font("Helvetica", "", 10)
        pdf.multi_cell(width, 4.6, _latin(text))

    def bullet(text: str) -> None:
        pdf.set_x(pdf.l_margin)
        pdf.set_font("Helvetica", "", 10)
        pdf.cell(4, 4.6, _latin("-"))
        pdf.multi_cell(width - 4, 4.6, _latin(text))

    name_block()

    heading("Profile")
    body(cv.profile)

    heading("Professional Experience")
    for job in cv.experience:
        pdf.ln(1.5)
        pdf.set_x(pdf.l_margin)
        pdf.set_font("Helvetica", "I", 9)
        date_w = pdf.get_string_width(_latin(job.dates)) + 2
        pdf.set_font("Helvetica", "B", 10.5)
        pdf.cell(width - date_w, 5, _latin(job.title))
        pdf.set_font("Helvetica", "I", 9)
        pdf.cell(date_w, 5, _latin(job.dates), align="R", new_x="LMARGIN", new_y="NEXT")
        pdf.ln(0.5)
        for b in job.bullets:
            bullet(b)

    heading("Education")
    for school, dates in cv.education:
        pdf.set_x(pdf.l_margin)
        pdf.set_font("Helvetica", "B", 10)
        pdf.multi_cell(width, 4.6, _latin(f"{school}  -  {dates}"))

    heading("Technical Skills")
    for cat, items in cv.skills:
        pdf.set_font("Helvetica", "B", 10)
        pdf.write(4.8, _latin(cat + ": "))
        pdf.set_font("Helvetica", "", 10)
        pdf.write(4.8, _latin(items))
        pdf.ln(5.2)

    heading("Selected AI Systems")
    for name, desc in cv.selected_systems:
        pdf.set_x(pdf.l_margin)
        pdf.set_font("Helvetica", "B", 10)
        pdf.write(4.8, _latin("- " + name))
        pdf.set_font("Helvetica", "", 10)
        pdf.write(4.8, _latin(" - " + desc))
        pdf.ln(5.2)

    heading("Certifications & Courses")
    body(cv.courses)

    heading("Additional")
    body(cv.additional)

    pdf.output(path)


# --- Markdown renderer ------------------------------------------------------

def render_md(cv: CVData, path: str) -> None:
    contact = " · ".join(f"[{t}]({u})" if u else t for t, u in cv.contact_segments)
    lines = [f"# {cv.name}", f"**{cv.role}**", "", contact, "", "## PROFILE", cv.profile, "",
             "## PROFESSIONAL EXPERIENCE"]
    for job in cv.experience:
        lines += ["", f"**{job.title}** · {job.dates}"]
        lines += [f"- {b}" for b in job.bullets]
    lines += ["", "## EDUCATION"]
    lines += [f"**{s}** · {d}  " for s, d in cv.education]
    lines += ["", "## TECHNICAL SKILLS"]
    lines += [f"- **{c}:** {i}" for c, i in cv.skills]
    lines += ["", "## SELECTED AI SYSTEMS"]
    lines += [f"- **{n}** — {d}" for n, d in cv.selected_systems]
    lines += ["", "## CERTIFICATIONS & COURSES", cv.courses, "", "## ADDITIONAL", cv.additional, ""]
    with open(path, "w", encoding="utf-8") as f:
        f.write("\n".join(lines))
